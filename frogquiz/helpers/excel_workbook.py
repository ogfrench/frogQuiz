# SPDX-FileCopyrightText: 2026 frogQuiz contributors
#
# SPDX-License-Identifier: MPL-2.0

"""Text-only Excel validation. Parsing never reads or writes application state."""

import re
from io import BytesIO
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError
from zlib import error as ZlibError

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
# Matches the editor's user-facing limit. Never silently truncate a larger sheet.
MAX_IMPORT_QUESTIONS = 50
MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = [
    "Question",
    "1st Answer",
    "2nd Answer",
    "3rd Answer",
    "4th Answer",
    "Time Limit (max. 120)",
    "Correct Answers",
]


def parse_excel(data: bytes, *, title_limit=100, description_limit=500, question_limit=250, answer_limit=100):
    preview = {"title": "", "description": "", "questions": [], "errors": []}

    def error(row, column, message):
        preview["errors"].append({"row": row, "column": column, "message": message})

    if len(data) > MAX_FILE_BYTES:
        error(None, None, "The file must be at most 2 MB.")
        return preview
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 200 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                error(None, None, "The workbook is too large when unpacked. Use the text-only template.")
                return preview
            if any(entry.flag_bits & 1 for entry in entries):
                error(None, None, "Password-protected workbooks cannot be imported.")
                return preview
            for entry in entries:
                if entry.filename.endswith((".xml", ".rels")):
                    # UTF-16/32 declarations must be refused just like UTF-8 ones.
                    xml = archive.read(entry).replace(b"\x00", b"").upper()
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        error(None, None, "The workbook contains unsupported XML declarations. Use the template.")
                        return preview
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        ValueError,
        TypeError,
        OSError,
        ParseError,
        SyntaxError,
        ZlibError,
        NotImplementedError,
    ):
        error(None, None, "This file could not be read as an .xlsx workbook. Use the template.")
        return preview

    try:
        sheet = workbook["Main"] if "Main" in workbook.sheetnames else workbook.active
        if sheet is None:
            error(None, None, "The workbook needs a worksheet.")
            return preview
        # The export and the original import template use this exact layout.
        rows = list(sheet.iter_rows(min_row=1, max_row=13 + MAX_IMPORT_QUESTIONS + 1, max_col=14))

        def text(cell, row, column, limit, required=False):
            if cell.data_type == "f":
                error(row, column, "Use a value, not a formula.")
                return ""
            value = "" if cell.value is None else str(cell.value).strip()
            if required and not value:
                error(row, column, "This cell is required.")
            elif len(value) > limit:
                error(row, column, f"Use at most {limit} characters.")
            return value

        preview["title"] = text(rows[4][2], 5, "Title", title_limit, required=True)
        if preview["title"] and len(preview["title"]) < 3:
            error(5, "Title", "Use at least 3 characters.")
        preview["description"] = text(rows[5][2], 6, "Description", description_limit)
        headers = [str(cell.value or "").strip() for cell in rows[12][1:]]
        time_index = next((i for i, name in enumerate(headers) if name.startswith("Time")), None)
        if (
            headers[0] != "Question"
            or time_index is None
            or not 3 <= time_index <= 11
            or headers[time_index + 1].lower() != "correct answers"
        ):
            error(13, None, "Keep the template's Question, answer, Time, and Correct Answers columns.")
            return preview
        answer_count = time_index - 1
        for index, header in enumerate(headers[1:time_index], start=1):
            match = re.fullmatch(r"(?:(\d+)(?:st|nd|rd|th)? Answer|Answer (\d+))", header)
            if match is None or int(match[1] or match[2]) != index:
                error(13, None, "Keep the template's answer column headings in order.")
                return preview
        if sheet.max_row is not None and sheet.max_row > 13 + MAX_IMPORT_QUESTIONS:
            error(None, None, f"Use at most {MAX_IMPORT_QUESTIONS} question rows, starting at row 14.")

        found = False
        for number, cells in enumerate(rows[13 : 13 + MAX_IMPORT_QUESTIONS], start=14):
            cells = cells[1 : time_index + 3]
            if all(cell.value is None or cell.value == "" for cell in cells):
                continue
            found = True
            before = len(preview["errors"])
            question = text(cells[0], number, "Question", question_limit, required=True)
            answers = [text(cells[i + 1], number, f"Answer {i + 1}", answer_limit) for i in range(answer_count)]
            if sum(bool(answer) for answer in answers) < 2:
                error(number, "Answers", "Use at least 2 nonempty answers.")
            timer = cells[time_index]
            try:
                seconds = float(timer.value)
                if (
                    timer.data_type == "f"
                    or isinstance(timer.value, bool)
                    or not seconds.is_integer()
                    or not 1 <= seconds <= 120
                ):
                    raise ValueError
                seconds = int(seconds)
            except (TypeError, ValueError, OverflowError):
                error(number, "Time", "Use a whole number from 1 to 120.")
                seconds = 20
            correct = text(cells[time_index + 1], number, "Correct answers", 100, required=True)
            # Excel may store a single selection as a number rather than text.
            if re.fullmatch(r"\d+\.0", correct):
                correct = correct[:-2]
            selected = set()
            if correct and not re.fullmatch(r"\d+(?:\s*,\s*\d+)*", correct):
                error(number, "Correct answers", "Use answer numbers separated by commas, such as 1,3.")
            elif correct:
                selected = {int(part.strip()) for part in correct.split(",")}
                for index in sorted(selected):
                    if not 1 <= index <= answer_count or not answers[index - 1]:
                        error(number, "Correct answers", f"Answer {index} is missing or empty.")
            if len(preview["errors"]) == before:
                preview["questions"].append(
                    {
                        "row": number,
                        "question": question,
                        "type": "multiple" if len(selected) > 1 else "single",
                        "time": seconds,
                        "answers": [
                            {"text": answer, "right": index in selected}
                            for index, answer in enumerate(answers, start=1)
                            if answer
                        ],
                    }
                )
        if not found:
            error(None, "Questions", "Add at least one question starting at row 14.")
        return preview
    except (
        ValueError,
        TypeError,
        KeyError,
        IndexError,
        OSError,
        BadZipFile,
        ParseError,
        SyntaxError,
        ZlibError,
        NotImplementedError,
    ):
        # A structurally broken workbook is a user-correctable file error.
        error(None, None, "The worksheet could not be read. Save a new copy of the template.")
        return preview
    finally:
        workbook.close()


def excel_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["B2"] = "frogQuiz: create a quiz from Excel"
    sheet["B3"] = "Fill in the title, replace the sample questions, then save as .xlsx."
    sheet["B5"], sheet["C5"] = "Title", "Capital cities of Europe"
    sheet["B6"], sheet["C6"] = "Description (optional)", "A warm-up"
    sheet["B8"] = "Up to 50 questions. Question: 250 characters. Each answer: 100 characters."
    sheet["B9"] = "Correct answers: use column numbers, such as 1 or 1,3. Leave unused answers blank."
    sheet["B10"] = "Text only. No formulas or pictures. A confirmation creates a new quiz."
    for column, heading in enumerate(HEADERS, start=2):
        cell = sheet.cell(13, column, heading)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="27272A")
    for number, values in enumerate(
        [
            ["What is the capital of Portugal?", "Lisbon", "Porto", "Faro", None, 20, "1"],
            ["Which of these are capitals?", "Madrid", "Barcelona", "Rome", "Milan", 30, "1,3"],
        ],
        start=14,
    ):
        for column, value in enumerate(values, start=2):
            sheet.cell(number, column, value)
    for column in "CDEF":
        sheet.column_dimensions[column].width = 23
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["G"].width = 24
    sheet.column_dimensions["H"].width = 22
    sheet.freeze_panes = "C14"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
